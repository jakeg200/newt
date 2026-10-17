"""Turn uploaded documents into citable passages.

A passage is a few paragraphs (roughly 200–1200 characters) so a handbook rule can point at the
exact text it came from. Headings and short lines are merged into the paragraph that follows.
"""

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

MAX_BYTES = 20 * 1024 * 1024
MIN_CHARS = 200
MAX_CHARS = 1200

MIME_BY_SUFFIX = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


class ExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class Block:
    text: str
    page: int | None = None


def mime_for(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in MIME_BY_SUFFIX:
        raise ExtractionError(f"unsupported file type {suffix or '(none)'}; use PDF, Word (.docx), .txt or .md")
    return MIME_BY_SUFFIX[suffix]


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _split_paragraphs(text: str, page: int | None = None) -> list[Block]:
    return [Block(p, page) for p in (_clean(p) for p in re.split(r"\n\s*\n", text)) if p]


def _blocks_from_pdf(data: bytes) -> list[Block]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise ExtractionError(f"could not read PDF ({exc})") from exc
    return [b for n, text in enumerate(pages, start=1) for b in _split_paragraphs(text, page=n)]


def _blocks_from_docx(data: bytes) -> list[Block]:
    import zipfile

    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("could not read Word document") from exc
    blocks = [Block(_clean(p.text)) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [_clean(c.text) for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(Block(" | ".join(cells)))
    return blocks


def _blocks_from_text(data: bytes) -> list[Block]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    return _split_paragraphs(text)


def _split_long(block: Block) -> list[Block]:
    """Split an over-long block on sentence boundaries."""
    if len(block.text) <= MAX_CHARS:
        return [block]
    sentences = re.split(r"(?<=[.!?;:])\s+", block.text)
    out: list[Block] = []
    current = ""
    for sentence in sentences:
        while len(sentence) > MAX_CHARS:  # a run-on with no punctuation
            if current:
                out.append(Block(current, block.page))
                current = ""
            out.append(Block(sentence[:MAX_CHARS], block.page))
            sentence = sentence[MAX_CHARS:]
        if current and len(current) + 1 + len(sentence) > MAX_CHARS:
            out.append(Block(current, block.page))
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        out.append(Block(current, block.page))
    return out


def chunk(blocks: list[Block]) -> list[Block]:
    """Merge short blocks forward and split long ones, keeping the page of the first block."""
    passages: list[Block] = []
    pending: Block | None = None
    for block in (piece for b in blocks for piece in _split_long(b)):
        if pending is None:
            pending = block
        elif len(pending.text) + 2 + len(block.text) <= MAX_CHARS:
            pending = Block(f"{pending.text}\n\n{block.text}", pending.page)
        else:
            passages.append(pending)
            pending = block
        if len(pending.text) >= MIN_CHARS:
            passages.append(pending)
            pending = None
    if pending is not None:
        if passages and len(passages[-1].text) + 2 + len(pending.text) <= MAX_CHARS:
            passages[-1] = Block(f"{passages[-1].text}\n\n{pending.text}", passages[-1].page)
        else:
            passages.append(pending)
    return passages


def extract_passages(filename: str, data: bytes) -> list[Block]:
    if len(data) > MAX_BYTES:
        raise ExtractionError(f"file is larger than {MAX_BYTES // (1024 * 1024)} MB")
    mime = mime_for(filename)
    if mime == "application/pdf":
        blocks = _blocks_from_pdf(data)
    elif mime.startswith("application/vnd.openxmlformats"):
        blocks = _blocks_from_docx(data)
    else:
        blocks = _blocks_from_text(data)
    passages = chunk(blocks)
    if not passages:
        raise ExtractionError("no text found (a scanned PDF needs OCR, which isn't supported yet)")
    return passages
